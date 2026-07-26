"""Build the TryOnReady Word/PDF source guide.

Design preset: compact_reference_guide
Header template: customer_pack

The PDF is emitted by the official document-rendering workflow, not by this
script. This script intentionally contains no provider credential.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "judge" / "TryOnReady-Judge-and-Use-Case-Guide.docx"
ASSETS = ROOT / "samples" / "synthetic"

BLUE = "2E74B5"
DARK = "2A201D"
RUST = "A84E37"
PALE_BLUE = "EAF3F9"
PALE_GREEN = "EAF4E7"
PALE_SAND = "F8F2EA"
GRAY = "666666"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_width(cell, width_inches: float) -> None:
    width = Inches(width_inches)
    cell.width = width
    properties = cell._tc.get_or_add_tcPr()
    tc_width = properties.find(qn("w:tcW"))
    if tc_width is None:
        tc_width = OxmlElement("w:tcW")
        properties.append(tc_width)
    tc_width.set(qn("w:w"), str(int(width.twips)))
    tc_width.set(qn("w:type"), "dxa")


def set_table_width(table, width_inches: float) -> None:
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(int(Inches(width_inches).twips)))
    table_width.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    row_properties.append(repeat)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(9.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)


def style_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(DARK)
    title.paragraph_format.space_after = Pt(12)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = RGBColor.from_string(RUST)
    subtitle.paragraph_format.space_after = Pt(14)

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, RUST, 10, 5),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left = paragraph.add_run("TryOnReady | Judge and Small-Business Guide")
    left.bold = True
    left.font.name = "Calibri"
    left.font.size = Pt(9)
    left.font.color.rgb = RGBColor.from_string(BLUE)
    right = paragraph.add_run("\tVerified demo journey")
    right.font.name = "Calibri"
    right.font.size = Pt(9)
    right.font.color.rgb = RGBColor.from_string(GRAY)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.4))

    footer = section.footer
    add_page_field(footer.paragraphs[0])


def add_kicker(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(RUST)


def add_callout(document: Document, title: str, body: str, fill: str = PALE_BLUE) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, 6.7)
    cell = table.cell(0, 0)
    set_cell_width(cell, 6.7)
    set_cell_shading(cell, fill)
    cell.margin_top = Inches(0.12)
    cell.margin_bottom = Inches(0.12)
    cell.margin_left = Inches(0.15)
    cell.margin_right = Inches(0.15)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(3)
    heading = paragraph.add_run(title)
    heading.bold = True
    heading.font.name = "Calibri"
    heading.font.size = Pt(11)
    heading.font.color.rgb = RGBColor.from_string(DARK)
    paragraph = cell.add_paragraph(body)
    paragraph.paragraph_format.space_after = Pt(0)


def add_numbered(document: Document, steps: list[str]) -> None:
    numbering = document.part.numbering_part.element
    style_num_id = int(
        document.styles["List Number"]
        ._element.pPr.numPr.numId.val
    )
    source_num = numbering.find(
        f"{qn('w:num')}[@{qn('w:numId')}='{style_num_id}']"
    )
    if source_num is None:
        raise RuntimeError("The List Number numbering definition is missing.")

    abstract_id = source_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    existing_ids = [
        int(element.get(qn("w:numId")))
        for element in numbering.findall(qn("w:num"))
    ]
    num_id = max(existing_ids, default=0) + 1

    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    number.append(abstract)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    number.append(level_override)
    numbering.append(number)

    for step in steps:
        paragraph = document.add_paragraph(step, style="List Number")
        paragraph.paragraph_format.keep_together = True
        paragraph_properties = paragraph._p.get_or_add_pPr()
        numbering_properties = paragraph_properties.find(qn("w:numPr"))
        if numbering_properties is None:
            numbering_properties = OxmlElement("w:numPr")
            paragraph_properties.append(numbering_properties)
        level = OxmlElement("w:ilvl")
        level.set(qn("w:val"), "0")
        direct_num_id = OxmlElement("w:numId")
        direct_num_id.set(qn("w:val"), str(num_id))
        numbering_properties.append(level)
        numbering_properties.append(direct_num_id)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(item, style="List Bullet")
        paragraph.paragraph_format.keep_together = True


def add_picture(document: Document, path: Path, alt: str, width: float) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    inline = document.inline_shapes[-1]._inline
    inline.docPr.set("descr", alt)


def add_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GRAY)


def add_credentials(document: Document) -> None:
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, 6.7)
    widths = [1.1, 2.0, 1.75, 1.85]
    headers = ["Role", "Username", "Password", "Purpose"]
    for index, (header, width) in enumerate(zip(headers, widths)):
        set_cell_width(table.rows[0].cells[index], width)
        set_cell_shading(table.rows[0].cells[index], BLUE)
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
    set_repeat_header(table.rows[0])

    rows = [
        (
            "Retailer",
            "retailer@tryonready.demo",
            "RetailerDemo!2026",
            "Apply, add garments, view totals",
        ),
        (
            "Administrator",
            "admin@tryonready.demo",
            "AdminDemo!2026",
            "Approve boutique and garment",
        ),
        ("Guest", "None", "None", "Consent, upload, view result"),
    ]
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, (value, width) in enumerate(zip(values, widths)):
            set_cell_width(cells[index], width)
            if row_index % 2 == 0:
                set_cell_shading(cells[index], "F5F8FA")
            set_cell_text(cells[index], value)
    document.add_paragraph()


def add_verification_table(document: Document) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, 6.7)
    for index, (header, width) in enumerate((("Check", 4.8), ("Result", 1.9))):
        set_cell_width(table.rows[0].cells[index], width)
        set_cell_shading(table.rows[0].cells[index], BLUE)
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
    set_repeat_header(table.rows[0])
    rows = [
        ("Release build", "0 warnings / 0 errors"),
        (".NET unit tests", "2 passed"),
        (".NET integration tests", "8 passed"),
        ("Frontend lint and TypeScript", "Passed"),
        ("Playwright desktop journey", "Passed"),
        ("Playwright mobile smoke", "Passed"),
        ("PostgreSQL-backed journey", "Passed"),
        ("Duplicate-request protection", "Passed"),
        ("API key exposed to browser", "False"),
        ("Isolated PostgreSQL", "Healthy"),
    ]
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, (value, width) in enumerate(zip(values, (4.8, 1.9))):
            set_cell_width(cells[index], width)
            if row_index % 2 == 0:
                set_cell_shading(cells[index], "F5F8FA")
            set_cell_text(cells[index], value, bold=index == 1)


def add_product(
    document: Document,
    *,
    name: str,
    sku: str,
    category: str,
    color: str,
    material: str,
    sizes: str,
    price: str,
    image: str,
    dimensions: str,
    file_size: str,
    sha256: str,
) -> None:
    document.add_heading(name, level=1)
    add_picture(
        document,
        ASSETS / "garments" / image,
        f"Original synthetic product photograph of the fictional {name}",
        3.0 if category == "Top" else 2.45,
    )
    add_caption(
        document,
        f"Authorized synthetic garment input. File: {image}",
    )
    details = [
        f"Product number: {sku}",
        f"Category: {category}",
        "Brand: Luna & Thread",
        f"Color: {color}",
        f"Material: {material}",
        f"Sizes: {sizes}",
        f"Demo price: {price}",
        f"Image: {dimensions}; {file_size}",
        f"SHA-256: {sha256}",
    ]
    add_bullets(document, details)
    add_callout(
        document,
        "Upload-ready",
        "The complete garment is visible on a neutral background with no model, "
        "hanger, logo, text, or third-party brand.",
        PALE_GREEN,
    )


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    style_document(document)
    document.core_properties.title = "TryOnReady Judge and Use-Case Guide"
    document.core_properties.subject = (
        "Lean end-to-end hackathon demonstration and synthetic catalog"
    )
    document.core_properties.author = "TryOnReady"
    document.core_properties.keywords = (
        "TryOnReady, YouCam, virtual try-on, judge guide, Luna & Thread"
    )

    add_kicker(document, "Private hackathon demonstration")
    document.add_heading("TryOnReady", level=0)
    document.add_paragraph(
        "Judge and Small-Business Use-Case Guide",
        style="Subtitle",
    )
    document.add_paragraph(
        "One continuous, easy-to-evaluate journey from boutique application "
        "to a secure virtual try-on result."
    )
    add_callout(
        document,
        "The complete story",
        "Elena Rivera submits Luna & Thread, adds a synthetic garment, an "
        "administrator approves it, and a guest customer consents and selects "
        "a synthetic photo. The secure ASP.NET server calls the configured "
        "provider and returns the result. The API key never reaches the browser "
        "or GitHub.",
        PALE_SAND,
    )
    document.add_paragraph()
    document.add_paragraph("Prepared: July 26, 2026")
    document.add_paragraph("Architecture: ASP.NET Core 10 + Next.js/TypeScript + EF Core 10 + PostgreSQL")
    document.add_paragraph("Demo boutique: Luna & Thread (fictional)")
    document.add_paragraph("Retailer persona: Elena Rivera (fictional)")
    document.add_page_break()

    document.add_heading("Start here: roles and access", level=1)
    document.add_paragraph(
        "The retailer and administrator sign in. The guest customer does not "
        "need an account for this lean MVP."
    )
    add_credentials(document)
    add_callout(
        document,
        "Local URL",
        "Start the TryOnReady.Api https profile in Visual Studio, then open "
        "https://localhost:7090/. The HTTP profile is http://localhost:5090/.",
    )
    document.add_heading("The 1-3 minute judge path", level=2)
    add_numbered(
        document,
        [
            "Retailer signs in and submits Luna & Thread.",
            "Retailer uploads the Moonlight Blazer and passes image readiness.",
            "Administrator signs in and approves the boutique and garment.",
            "Guest selects the approved garment, consents, and uploads Marisol Lopez's synthetic photo.",
            "The server creates and tracks the provider task; the guest sees the result.",
            "A repeated click is stopped as a duplicate; the dashboard shows aggregate totals.",
        ],
    )
    add_callout(
        document,
        "Never enter the YouCam key in the browser",
        "The provider key is a server-side secret. It is intentionally absent "
        "from this guide and every repository file.",
        PALE_GREEN,
    )
    document.add_page_break()

    document.add_heading("Retailer journey: Elena Rivera", level=1)
    document.add_heading("1. Submit Luna & Thread", level=2)
    add_numbered(
        document,
        [
            "Open Judge Sign In and choose Retailer.",
            "Enter the retailer credentials exactly as shown.",
            "Keep the prepared Luna & Thread and Elena Rivera values.",
            "Check the image-rights statement.",
            "Select Submit application and confirm Application submitted.",
        ],
    )
    document.add_heading("2. Add the Moonlight Blazer", level=2)
    add_numbered(
        document,
        [
            "Continue to Product Readiness.",
            "Keep product number SYN-BLZ-001 and category Top.",
            "Choose samples/synthetic/garments/moonlight-blazer.png.",
            "Confirm the image reads 1254 x 1254 pixels.",
            "Select Save garment and check image.",
            "Confirm the product is waiting for administrator approval.",
        ],
    )
    add_picture(
        document,
        ASSETS / "garments" / "moonlight-blazer.png",
        "Original synthetic Moonlight Blazer product photograph",
        2.75,
    )
    add_caption(
        document,
        "Moonlight Blazer | SYN-BLZ-001 | Top | Terracotta | USD 89",
    )
    document.add_page_break()

    document.add_heading("Administrator and guest journey", level=1)
    document.add_heading("3. Human approval", level=2)
    add_numbered(
        document,
        [
            "Return to Judge Sign In and choose Administrator.",
            "Enter the administrator credentials exactly as shown.",
            "Approve Luna & Thread.",
            "Inspect the Moonlight Blazer image and catalog facts.",
            "Approve the product.",
            "Continue to Consumer Try-On.",
        ],
    )
    document.add_heading("4. Guest customer try-on", level=2)
    add_numbered(
        document,
        [
            "Confirm Moonlight Blazer is selected.",
            "Confirm the page says API key in browser: Never.",
            "Choose Marisol Lopez's synthetic customer image.",
            "Check the consent statement.",
            "Select Generate virtual try-on once.",
            "Wait for Succeeded and confirm the generated result appears.",
            "Submit the unchanged request again and confirm Duplicate request prevented.",
            "Open the results dashboard and confirm aggregate totals.",
        ],
    )
    add_picture(
        document,
        ASSETS / "customers" / "marisol-lopez-source.png",
        "Fictional adult customer Marisol Lopez, authorized synthetic full-body source image",
        1.7,
    )
    add_caption(
        document,
        "Marisol Lopez — fictional adult synthetic customer source, 864 x 1821",
    )
    document.add_page_break()

    document.add_heading("Fictional customer alternatives", level=1)
    document.add_paragraph(
        "Both people are fictional adults created specifically for this demo. "
        "They are not real customers. Use only these authorized files for the "
        "controlled live validation."
    )
    document.add_heading("Marisol Lopez", level=2)
    add_picture(
        document,
        ASSETS / "customers" / "marisol-lopez-source.png",
        "Fictional adult customer Marisol Lopez synthetic source",
        1.45,
    )
    add_caption(
        document,
        "customers/marisol-lopez-source.png | 864 x 1821 | full-body source",
    )
    document.add_heading("Danielle Smith", level=2)
    add_picture(
        document,
        ASSETS / "customers" / "danielle-smith-source.png",
        "Fictional adult customer Danielle Smith synthetic source",
        1.45,
    )
    add_caption(
        document,
        "customers/danielle-smith-source.png | 864 x 1821 | full-body source",
    )
    document.add_page_break()

    add_product(
        document,
        name="Moonlight Blazer",
        sku="SYN-BLZ-001",
        category="Top",
        color="Terracotta",
        material="Cotton-blend suiting",
        sizes="XS-XL",
        price="USD 89.00",
        image="moonlight-blazer.png",
        dimensions="1254 x 1254",
        file_size="1,788,350 bytes",
        sha256="8f02878034877b38cc1e9f8372bfd145819e6ffff74f5e295709e3af8f5fd907",
    )
    document.add_page_break()

    add_product(
        document,
        name="Harbor Sage Blouse",
        sku="SYN-TOP-002",
        category="Top",
        color="Sage",
        material="Cotton-linen blend",
        sizes="XS-XXL",
        price="USD 64.00",
        image="harbor-sage-blouse.png",
        dimensions="1254 x 1254",
        file_size="1,608,884 bytes",
        sha256="8fd8adb4386ec064f9f8f621469aa685410abe7c7b7da62bb9e9fa25030a5bdc",
    )
    document.add_page_break()

    add_product(
        document,
        name="Midnight Wrap Dress",
        sku="SYN-DRS-003",
        category="Full-body outfit",
        color="Deep navy",
        material="Woven rayon blend",
        sizes="XS-XL",
        price="USD 118.00",
        image="midnight-wrap-dress.png",
        dimensions="1024 x 1536",
        file_size="2,131,211 bytes",
        sha256="4173e3974e3423db5c508c864cb0ac854a0d6d49d007af481070c58f9b6cf16e",
    )
    document.add_page_break()

    document.add_heading("Verification evidence", level=1)
    document.add_paragraph(
        "These results were recorded on July 26, 2026 using provider simulation "
        "for the automated journey. Simulation proves application behavior "
        "without spending a YouCam unit."
    )
    add_verification_table(document)
    document.add_paragraph()
    add_callout(
        document,
        "YouCam account ready",
        "The YouCam account balance screenshot confirms 1,040 bonus units. A "
        "controlled live call must use the same authorized assets, be submitted "
        "once, and be recorded in docs/IMPLEMENTATION-LOG.md.",
        PALE_GREEN,
    )
    document.add_heading("Security and privacy facts", level=2)
    add_bullets(
        document,
        [
            "The YouCam API key is stored only in server-side secrets.",
            "No API key, redemption code, or signed upload URL is committed to GitHub.",
            "Consumer images are private and outside the public web root.",
            "The retailer sees totals and outcomes, not customer photographs.",
            "TryOnReady PostgreSQL is isolated on local port 55432.",
            "The existing Duevara and CarPermit PostgreSQL container is never changed.",
        ],
    )
    document.add_page_break()

    document.add_heading("Fast troubleshooting and scope", level=1)
    document.add_heading("If a step does not work", level=2)
    add_bullets(
        document,
        [
            "Applications will not load: sign in as Retailer again.",
            "Admin queues will not load: sign in as Administrator.",
            "No guest garment appears: approve both the boutique and product.",
            "Image rejected: use an unchanged repository synthetic file.",
            "Processing does not finish: check the API server log for provider status.",
            "Live provider disabled: complete simulation first, then follow docs/YOUCAM-SETUP.md.",
        ],
    )
    document.add_heading("Lean MVP boundary", level=2)
    document.add_paragraph(
        "The retailer pays for the service and offers virtual try-on to guests "
        "as a courtesy and convenience. A later phase may let a customer create "
        "an optional account, save favorite garments, signal purchase interest, "
        "and return to prior choices. Customer accounts, retailer subscriptions, "
        "checkout, and marketing automation are intentionally outside the "
        "hackathon MVP."
    )
    add_callout(
        document,
        "Source of truth",
        "For exact hashes and provenance, see samples/synthetic/ASSET-METADATA.md. "
        "For retention rules, see docs/DATA-RETENTION.md. For the provider "
        "secret and controlled live test, see docs/YOUCAM-SETUP.md.",
    )

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
